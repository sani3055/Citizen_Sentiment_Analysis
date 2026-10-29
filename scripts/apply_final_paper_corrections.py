"""Apply the October final-review corrections to the last revised manuscript."""

from copy import deepcopy
import json
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_journal_manuscript import format_math_notation

SOURCE = ROOT / "TeamNo_16_23BCT0029_MINOR_CORRECTIONS_FINAL.docx"
OUTPUT = ROOT / "TeamNo_16_23BCT0029_FINAL_CORRECTIONS.docx"


def replace_paragraph(paragraph, text):
    original = paragraph.runs[0] if paragraph.runs else None
    properties = deepcopy(original._r.rPr) if original is not None and original._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if properties is not None:
        run._r.insert(0, properties)
    run.font.superscript = False
    run.font.subscript = False
    format_math_notation(paragraph)


def insert_paragraph_after(anchor, text, template):
    element = OxmlElement("w:p")
    anchor.addnext(element)
    paragraph = Paragraph(element, template._parent)
    if template._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    format_math_notation(paragraph)
    return paragraph


def new_table_after(doc, anchor, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = Inches(width)
    for j, text in enumerate(headers):
        table.cell(0, j).text = text
    for row in rows:
        for cell, text, width in zip(table.add_row().cells, row, widths):
            cell.width = Inches(width)
            cell.text = str(text)
    for i, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "DFE9F2" if i == 0 else ("F5F7F9" if i % 2 == 0 else "FFFFFF"))
            tcpr.append(shd)
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                edge = OxmlElement("w:" + side)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    edge.set(qn("w:" + key), value)
                borders.append(edge)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                edge = OxmlElement("w:" + side)
                edge.set(qn("w:w"), "75")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            tcpr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
                    run.bold = i == 0
                format_math_notation(paragraph)
    anchor.addnext(table._tbl)
    return table


def main():
    doc = Document(SOURCE)
    original = list(doc.paragraphs)
    tables = list(doc.tables)
    selected = json.loads((ROOT / "reports/metrics/selected_feature_results.json").read_text())
    ranking = json.loads((ROOT / "reports/metrics/priority_ranking_metrics.json").read_text())

    prefix = original[5].text.split("On 994 held-out")[0]
    replace_paragraph(original[5], prefix +
        "On 994 NYC-2024 test cells, the all-family LightGBM delay model changed R2 from 0.5995 to 0.4730 (Delta R2 = -0.1264). This drop does not mean that every language family is unhelpful: topics, ABSA, and RSAT each improved delay prediction separately, while VADER reduced R2 by 0.2350. An exploratory model selected using 2023 validation data retained ABSA and RSAT and reached delay R2 0.6228 (Delta R2 = +0.0233; paired absolute-error p = 0.1087). Its bootstrap interval includes zero. For repeat rate, the all-family change was +0.0017, but validation selected structured features alone. Last-value persistence remained stronger at R2 0.8839 for delay and 0.9469 for repeat rate. On the duplicate-controlled six-class ABSA test, DistilBERT macro-F1 was 0.1191, below the rule baseline at 0.1668. The dashboard recovered only 1 of the observed top 10 cells; a separate training-derived burden cutoff yielded 9 of 10 relevant cells under a broader definition. These exploratory NYC-2024 results identify feature-specific effects without establishing improvement over persistence or cross-city generalisation.")

    replace_paragraph(original[39], original[39].text.replace("Table 3.", "Table 2."))
    replace_paragraph(original[45], original[45].text.replace("Table 2.", "Table 3."))
    replace_paragraph(original[35], original[35].text.replace("R2_augmented - R2_baseline", "augmented R2 - baseline R2").replace("f0", "f_0").replace("f1", "f_1"))
    replace_paragraph(original[50], original[50].text +
        " For the reported macro-F1, the Negative, Neutral, and Positive F1 scores are averaged with equal weight and zero_division = 0. The Neutral class has zero RSAT support; it is nevertheless retained in the declared three-label metric.")
    replace_paragraph(original[63], original[63].text.replace("994 untouched cells", "994 held-out cells").replace(
        "The dashboard ranking is evaluated separately with Precision@k, lift over random, and NDCG@k.",
        "The dashboard ranking is evaluated separately using raw hits, Precision@k, and NDCG@k. Original top-k relevance changes with k, so its lift is only Precision@k multiplied by N/k and is not independent evidence of utility."))
    methods = insert_paragraph_after(original[63]._p,
        "The final-review extension evaluates all eight subsets of topics, ABSA, and RSAT, including structured features alone; VADER is excluded from this candidate set. Each candidate is fitted on the 578 training cells and selected by R2 on the 313 validation cells, with fewer features preferred on an exact tie. The chosen subset is then refitted on train plus validation and evaluated on the same 994 NYC-2024 test cells. Five matched seeds and 2,000 paired cell-bootstrap resamples quantify uncertainty. Because the review and earlier ablations already examined 2024 outcomes, this is an exploratory refinement, not a fresh confirmatory test. No 2024 score is used to choose the subset within this extension.", original[70])
    insert_paragraph_after(methods._p,
        "A second ranking evaluation fixes relevance independently of k. Delay and repeat outcomes are converted to percentiles using February-August 2023 training distributions and combined as 55% delay plus 45% repeat. The cutoff is the 80th percentile of training burden scores (61.9031 on a 0-100 scale), frozen before applying the score to 2024. This yields 202 relevant test cells out of 994 (20.32%). Lift is then Precision@k divided by this same prevalence at every k. This exploratory cutoff is a burden proxy, not a validated municipal service standard.", original[70])
    replace_paragraph(original[67], original[67].text.replace("model artifacts, and a 20-test pytest suite", "analysis scripts, and a 24-test pytest suite") +
        " The public source repository is https://github.com/sani3055/Citizen_Sentiment_Analysis. The updated compact archive includes the processed cell data needed to rerun the selected-feature and ranking analyses; large raw request files and model weights are not redistributed in it.")
    replace_paragraph(original[70], original[70].text.replace(
        "Its macro-F1 was 0.2750, showing that performance was not balanced across classes.",
        "Its three-label macro-F1 was 0.2750: (0.8138 Negative + 0.0000 Neutral + 0.0113 Positive) / 3, rounded. Neutral has zero ground-truth support and receives zero F1 under the specified metric, so this score is not a two-class macro-F1."))
    replace_paragraph(original[73], original[73].text.replace(
        "With six classes, uniform random macro-F1 is approximately 0.167, so these results are at or below chance level. ",
        "Both learned models therefore underperformed the rule baseline. A six-class uniform-guess accuracy of 1/6 should not be confused with an established macro-F1 chance benchmark on this imbalanced test set. "))
    # Add the missing macro-F1 row to the existing NLP table without changing its other values.
    row = deepcopy(tables[4].rows[1]._tr)
    tables[4].rows[2]._tr.addnext(row)
    for cell, value in zip(tables[4].rows[3].cells, ["VADER", "1,747 cells; three-label average", "Macro-F1", "0.2750"]):
        cell.text = value
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)

    replace_paragraph(original[81], original[81].text.replace("untouched 2024", "held-out 2024"))
    replace_paragraph(original[84], original[84].text.replace("Across learned model families", "Within the original all-family comparisons") +
        " The family ablations are more specific than this all-family comparison: topics (+0.0122), ABSA (+0.0425), and RSAT (+0.0106) each improve delay R2, whereas VADER alone reduces it by 0.2350. Combining all families is therefore a poor feature-bundle result, not evidence that every civic-language signal lacks value.")

    cap = insert_paragraph_after(tables[7]._tbl,
        "Table 9. Exploratory validation-selected LightGBM models without VADER on the previously examined NYC-2024 horizon. Selection uses only the 2023 validation scores; seed 42 results are shown.", original[85])
    rows = []
    for label, target in [("Delay (hours)", "mean_resolution_delay_hours"), ("Repeat rate (14d)", "repeat_request_rate_14d")]:
        result = selected["results_by_target"][target]
        primary = result["primary_seed_42"]
        rows.append([label, " + ".join(result["selected_families"]) or "None", f"{result['selected_validation_metrics']['r2']:.4f}",
                     f"{primary['selected']['r2']:.4f}", f"{primary['delta_r2']:+.4f}", f"{primary['selected']['mae']:.4f}"])
    new = new_table_after(doc, cap._p, ["Target", "Added families", "Validation R2", "2024 R2", "Delta R2", "MAE"], rows, [.95, 1.15, 1.05, 1.0, 1.0, 1.1])
    para = insert_paragraph_after(new._tbl,
        "For delay, validation selected ABSA + RSAT and gave R2 0.6228, compared with 0.5995 for structured LightGBM (Delta +0.0233, paired absolute-error p = 0.1087). The all-non-VADER reference, using topics + ABSA + RSAT together, gave R2 0.6115. For repeat rate, validation preferred structured features alone (R2 0.9224), so no extra NLP family was retained. The corresponding all-non-VADER reference was 0.9240. Neither selected model surpassed the last-value baseline.", original[70])
    insert_paragraph_after(para._p,
        "Across five matched seeds, the selected delay gain was +0.0214 ± 0.0068 Delta R2. A 2,000-resample paired cell-bootstrap interval for seed-mean predictions was [-0.0263, +0.0762], which includes zero. Repeat-rate differences are exactly zero because the selected model is the structured model; that identity is not evidence of precise equivalence between different models. These exploratory findings support further feature-specific testing, not a confirmed improvement.", original[70])
    replace_paragraph(original[87], original[87].text.replace("Table 9.", "Table 10."))
    replace_paragraph(original[91], original[91].text)
    replace_paragraph(original[92], original[92].text + " These original robustness results concern the all-family bundle, not the selected model reported in Table 9.")
    replace_paragraph(original[96], original[96].text + " This figure explains the original all-family model, not the post-review selected-feature model.")
    replace_paragraph(original[101],
        "Table 11. Strict top-k overlap on NYC-2024 cells. Relevance is the observed top k under the composite burden score and therefore changes with k. The lift column is a mechanical rescaling of Precision@k, not separate evidence of strong operational performance.")
    legacy = tables[9]
    legacy.cell(0, 2).text = "Lift (rescaled P@k)"
    legacy.add_column(Inches(.7))
    legacy.cell(0, 4).text = "Hits"
    for row, (k, metric) in zip(list(legacy.rows)[1:], ranking["metrics_by_k"].items()):
        row.cells[4].text = f"{metric['hits']}/{k}"
    legacy.autofit = False
    for column, width in zip(legacy.columns, [.45, 1.25, 1.8, 1.25, 1.2]):
        column.width = Inches(width)
    for row in legacy.rows:
        for cell, width in zip(row.cells, [.45, 1.25, 1.8, 1.25, 1.2]):
            cell.width = Inches(width)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
    fixedcap = insert_paragraph_after(legacy._tbl,
        "Table 12. Separate fixed-cutoff evaluation on the same NYC-2024 queue. The training-derived cutoff produces 202 relevant cells among 994 (20.32% prevalence), held constant across k. This broader definition is not directly interchangeable with Table 11.", original[101])
    fixedrows = [[k, f"{m['hits']}/{k}", f"{m['precision_at_k']:.3f}", f"{m['recall_at_k']:.3f}", f"{m['lift_over_random']:.2f}"]
                 for k, m in ranking["fixed_threshold"]["metrics_by_k"].items()]
    fixedtable = new_table_after(doc, fixedcap._p, ["k", "Hits", "Precision@k", "Recall@k", "Lift vs 20.32%"], fixedrows, [.5, .8, 1.65, 1.5, 1.8])
    insert_paragraph_after(fixedtable._tbl,
        "The strict overlap is weak at the smallest queue sizes: 1/10 at k = 10 and 3/20 at k = 20. Its original 9.94-fold lift at k = 10 follows directly from 0.10 × 994/10. Under the separate fixed-cutoff burden definition, 9/10 and 18/20 cells qualify, with lift 4.43 and recall only 4.46% and 8.91%, respectively. A broad high-burden screen can therefore look useful without correctly recovering the most extreme observed top-k cells. The two definitions must be reported together; neither is resident-level validation or evidence that a model-led queue beats operational staff or a persistence-ranked queue.", original[70])

    replace_paragraph(original[104],
        "The all-family LightGBM bundle worsened delay prediction (Delta R2 -0.1264), but its failure should not be generalised to every added family. Topics, ABSA, and RSAT each improve delay in separate ablations, while VADER sharply reduces performance. Excluding VADER and selecting families on 2023 validation retains ABSA + RSAT and gives an exploratory delay gain of +0.0233 over structured LightGBM. The gain is uncertain: paired absolute-error p = 0.1087 and the bootstrap interval spans zero. Persistence remains stronger for both targets. The contribution is therefore a feature-specific lesson about municipal NLP under privacy and boilerplate constraints, not a claim that language never helps or that it beats a temporal benchmark.")
    replace_paragraph(original[105], original[105].text +
        " The approximately 0.92 repeat-rate R2 is consistent with temporal persistence and address concentration rather than demonstrated modelling gain: last-value R2 is higher at 0.9469, and validation retains no added language or satisfaction family. These aggregate comparisons do not causally identify how much of R2 is due to address concentration.")
    replace_paragraph(original[109], original[109].text +
        " All operational performance claims are limited to the selected NYC services and boroughs in the January-December 2024 target horizon; no cross-city or later-year replication is supplied.")
    replace_paragraph(original[110], original[110].text +
        " The feature-selection extension and fixed-cutoff ranking check are post-review exploratory analyses on an already examined test period. Independent temporal and external-city confirmation remains necessary.")
    replace_paragraph(original[112],
        "This study combines 4,204,783 NYC 311 requests and 406,581 survey responses to evaluate one-month-ahead prediction at privacy-safe service-borough-month grain. On the NYC-2024 horizon, the all-family augmentation worsens delay prediction, largely alongside the harmful VADER family, while separate topic, ABSA, and RSAT additions show modest positive delay effects. Validation-selected ABSA + RSAT raises delay R2 from 0.5995 to 0.6228, but its uncertainty includes no gain and it does not beat persistence. No extra family is selected for repeat rate, whose high R2 must be interpreted against stronger last-value prediction and concentrated address recurrence. DistilBERT also underperforms the rule-based grievance classifier. The dashboard has weak strict top-10 recovery despite stronger results under a separately defined broad burden cutoff. These exploratory, NYC-specific findings identify signals worth testing and failure modes worth reporting; they do not establish a general NLP advantage, causal effects, or transfer to another city.")
    replace_paragraph(original[117],
        "Data availability. The study uses public NYC Open Data Socrata datasets erm2-nwe9 and 5ijn-vbdv. Raw records can be reconstructed with the documented paged queries. The public code repository is https://github.com/sani3055/Citizen_Sentiment_Analysis. Code, configuration, processed cell tables, metric JSON files, figures, and reproducibility instructions are packaged in the final-corrections archive. Zenodo deposition and assignment of a permanent DOI remain pending author action; no DOI is claimed in this version.")

    # Keep the source's author superscripts, equation runs, figures, and layout.
    # Adjust only obsolete descriptions of the now-reused hold-out horizon.
    for paragraph in doc.paragraphs:
        text = paragraph.text
        replacement = text.replace("untouched 2024", "held-out 2024").replace("keeps a later period untouched for testing", "holds out a later period for testing")
        if text != replacement:
            replace_paragraph(paragraph, replacement)
    # Newly added captions stay with their tables; table body rows may wrap but never clip.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Table "):
            paragraph.paragraph_format.keep_with_next = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
